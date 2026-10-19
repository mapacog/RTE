import argparse
import csv
import json
import sys
from pathlib import Path
from openpyxl import load_workbook
sys.stdout.reconfigure(encoding='utf-8')
root = Path(__file__).resolve().parents[1]
out = root / 'public' / 'data'
parser = argparse.ArgumentParser(description="Prepare interactive report tables from the Tables1_2_3 workbook.")
parser.add_argument("workbook", type=Path, help="Path to the source .xlsx workbook")
book_path = parser.parse_args().workbook
if not book_path.is_file():
    parser.error(f"Workbook not found: {book_path}")
out.mkdir(parents=True, exist_ok=True)
book = load_workbook(book_path, data_only=True)
tables = []
rows = list(book['Table 1'].iter_rows(values_only=True))
headers = ['Roadway class', 'Centerline miles (legacy)', '2018-2019', '2019-2020', '2020-2021', '2021-2022', '2022-2023 methodology break', 'OSM network miles', '2023-2024', '2024-2025']
tables.append({'id':'ffc','title':'VMT change by roadway class','sourceSheet':'Table 1','columns':headers,'formats':['text','number','percent','percent','percent','percent','text','number','percent','percent'],'rows':[list(r[:10]) for r in rows[2:] if r[0]],'note':'The 2022-2023 methodology break separates the legacy centerline network from the OSM network. No directly comparable change is reported for that interval.'})
for row in tables[0]['rows']:
    if row[0] == 'Major Arterial':
        row[0] = 'Principal Arterial'
atr=[]
state=None
for r in list(book['Table 2'].iter_rows(values_only=True))[2:]:
    if not r[1] or str(r[1]).strip() == 'ID': continue
    state = r[0] or state
    atr.append([state,*r[1:5]])
tables.append({'id':'atr','title':'Regional continuous traffic counters','sourceSheet':'Table 2','columns':['State','Counter ID','Route','Roadway class','Location'],'formats':['text']*5,'rows':atr,'note':'Automatic Traffic Recorders count vehicles throughout the year.'})
vmt=[]
for r in list(book['Table 3'].iter_rows(values_only=True))[1:]:
    for i in [0,2,4]:
        if isinstance(r[i], (int,float)): vmt.append([r[i],r[i+1],'Estimated' if r[i]>2025 else 'Measured'])
vmt.sort(key=lambda r:r[0])
tables.append({'id':'vmt','title':'Regional VMT by year','sourceSheet':'Table 3','columns':['Year','100M VMT (source label)','Series'],'formats':['year','number','text'],'rows':vmt,'note':'2026-2028 are estimates. The source column is labeled 100M VMT; that unit label needs confirmation against the annual VMT totals before release.'})
for t in tables:
    with (out / (t['id']+'.csv')).open('w',newline='',encoding='utf-8') as f:
        w=csv.writer(f);w.writerow(t['columns']);w.writerows(t['rows'])
(out/'reportTables.json').write_text(json.dumps(tables,indent=2),encoding='utf-8')

print("Prepared 3 interactive table sources from", book_path)
